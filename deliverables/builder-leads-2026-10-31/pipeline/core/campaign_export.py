# Merge validated emails back onto the lead list and write:
#  1) the master workbook (updated email columns), via write_xlsx.py
#  2) a Mission Control campaign file: first sheet = one row per unique, Valid address (the importer reads sheet 1
#     and detects the "Email" heading); second sheet = Risky addresses to review before sending.
import pandas as pd, re, json, os, csv
W = '/tmp/leads/work'; OUTD = '/tmp/leads/out'
df = pd.read_pickle(f'{W}/final_df.pkl'); df['abn_raw'] = df.ABN.str.replace(' ', '')
best = pd.read_pickle(f'{W}/v2/email_best.pkl')
newsite = {}
BUILDKW = re.compile(r'(build|construct|home|renovat|reno|carpent|project|develop|roof|kitchen|bathroom|joiner|concret|deck|patio|extension|maint|propert|design|struct|frame|brick|plaster|til(e|ing)|inspect|cabinet|pool|landscap|fenc|shed|steel|glaz|window|door|floor|paint|plumb|electric|civil|excavat|earthwork|demoli|carport|granny|house|dwelling|living|residen|estate|architect|draft)', re.I)
notb = json.load(open(f'{W}/v2/not_builders.json')) if os.path.exists(f'{W}/v2/not_builders.json') else {}
if os.path.exists(f'{W}/v3/not_builders2.json'): notb.update(json.load(open(f'{W}/v3/not_builders2.json')))
if os.path.exists(f'{W}/v4/not_builders3.json'): notb.update(json.load(open(f'{W}/v4/not_builders3.json')))
for a_, note_ in notb.items():
    df.loc[df.abn_raw == a_, 'Org Type'] = 'Non-builder licence holder (' + re.sub(r'(?i)^not (a )?builders?\s*[-–:]?\s*', '', note_)[:40].strip() + ')'
for f in ('out_site', 'out_disc', 'out_ind', 'out_ws', 'v3/out_site', 'v3/out_disc', 'v3/out_ind', 'v3/out_ws2', 'v4/out_cc', 'v4/out_cc_com', 'v4/out_ws3', 'v5/out_places'):
    p = f'{W}/v2/{f}.jsonl' if '/' not in f else f'{W}/{f}.jsonl'
    if os.path.exists(p):
        for l in open(p):
            j = json.loads(l)
            if f.endswith('out_ind') and j.get('website') and not ((j.get('web_conf') or '').startswith('Verified') or BUILDKW.search(j['website'])): continue
            if j.get('website') and j['abn'] not in newsite: newsite[j['abn']] = (j['website'], j.get('web_conf') or '', j.get('phones') or [])
def fmt_phone(p):
    if not p: return ''
    if p[:2] == '04' or p[:4] in ('1300', '1800'): return f'{p[:4]} {p[4:7]} {p[7:]}'
    return f'{p[:2]} {p[2:6]} {p[6:]}'
for i, r in df.iterrows():
    a = r.abn_raw
    if not r.Website and a in newsite:
        df.at[i, 'Website'] = newsite[a][0]; df.at[i, 'Website Match'] = newsite[a][1]
    if not r.Phone and a in newsite and newsite[a][2]:
        df.at[i, 'Phone'] = fmt_phone(newsite[a][2][0])
        if not r['Contact Source']: df.at[i, 'Contact Source'] = 'Company website'
    if a in best.index:
        b = best.loc[a]
        df.at[i, 'Email'] = b.email
        df.at[i, 'Email Status'] = b.status
        df.at[i, 'Email Type'] = ('Role inbox' if b.role else 'Named person') + ' – ' + {'own-domain': 'company domain', 'mailbox provider': 'Gmail/Bigpond-type mailbox', 'other-domain': 'other domain'}.get(b.kind, b.kind)
        df.at[i, 'Email Source'] = b.source
        df.at[i, 'Email Found On'] = b.page or ''
        df.at[i, 'Email Check'] = ('Grammar OK (Mission Control rules); domain mail server: ' + b.mx) + ('; ' + b.reasons if b.reasons else '')
        df.at[i, 'Alt Emails'] = b.alts or ''
        df.at[i, 'Email Shared By (leads)'] = str(int(b.shared_by)) if b.shared_by > 1 else ''
    else:
        df.at[i, 'Email'] = ''
for c in ['Email Status', 'Email Type', 'Email Source', 'Email Found On', 'Email Check', 'Alt Emails', 'Email Shared By (leads)']:
    df[c] = df[c].fillna('')
# Spam Act inferred-consent condition: the publishing page must not say unsolicited messages are unwanted
optf = json.load(open(f'{W}/v2/optout_flags.json')) if os.path.exists(f'{W}/v2/optout_flags.json') else {}
m_ = df.abn_raw.isin(optf.keys()) & (df.Email != '')
df.loc[m_, 'Email Status'] = 'Do not email'
df.loc[m_, 'Email Check'] = df.loc[m_, 'abn_raw'].map(lambda a: 'Publishing page states unsolicited email is not wanted: "...' + optf[a][:140] + '..."')
print('opt-out flagged', int(m_.sum()))
df['_contact'] = (df.Phone != '') | (df.Email != '')
# priority now favours leads that can actually be emailed
df['Priority'] = df.apply(lambda r: 'A' if r._tier <= 2 and r.Email and r['Email Status'] == 'Valid' else ('B' if (r._tier <= 2 or (r._tier == 3 and r.Email) or r['Master Builders Member']) else 'C') if r.Priority != 'C' or r.Email else 'C', axis=1)
cols = list(df.columns)
ei = cols.index('Email')
newc = ['Email Status', 'Email Type', 'Email Source', 'Email Found On', 'Email Check', 'Alt Emails', 'Email Shared By (leads)']
cols = [c for c in cols if c not in newc]
cols = cols[:ei + 1] + newc + cols[ei + 1:]
df = df[cols]
df.to_pickle(f'{W}/final_df.pkl')
# ---- campaign file ----
SUFX = re.compile(r'\b(PTY\.?|LTD\.?|LIMITED|PROPRIETARY|P/L)\b', re.I)
from rapidfuzz import fuzz
def _clean(nm):
    nm = re.sub(r'^The Trustee for (the )?', '', nm, flags=re.I)
    nm = re.sub(r'\s+', ' ', SUFX.sub('', nm)).strip(' .,&')
    nm = re.sub(r'\b(UNIT|FAMILY|DISCRETIONARY|HYBRID)?\s*TRUST$', '', nm, flags=re.I).strip(' .,&')
    return nm
def greet(r):
    # the name that best matches the business's own domain (website, else email domain); legal name otherwise
    names = [r.Company] + [x for x in r['Trading / Business Names'].split(' | ') if x]
    dom = (re.sub(r'^https?://(www\.)?', '', r.Website or '') or r.Email.split('@')[1]).split('.')[0].lower()
    if r['Email Type'].endswith('Gmail/Bigpond-type mailbox') and not r.Website: dom = r.Email.split('@')[0].lower()
    key = lambda n: re.sub(r'[^a-z0-9]', '', _clean(n).lower())
    scored = sorted(((fuzz.ratio(key(n), dom), i, n) for i, n in enumerate(names) if key(n)), reverse=True)
    best_n = scored[0][2] if scored and scored[0][0] >= 60 else r.Company
    nm = _clean(best_n)
    if r['_ind'] and best_n == r.Company: nm = ''  # never greet a sole trader by legal personal name pulled from a register
    if nm.isupper():
        nm = nm.title()
        nm = re.sub(r"\b(Wa|Nsw|Qld|Vic|Sa|Nt|Act|Tas|Nq|Fnq|Seq|Aus|Au)\b", lambda m: m.group(1).upper(), nm)
        nm = re.sub(r"(\w)'S\b", r"\1's", nm)
    return nm or 'there'
camp = df[(df.Email != '') & (~df['Org Type'].str.startswith('Non-builder'))].copy()
camp['Greeting Name'] = camp.apply(greet, axis=1)
camp['Lead Type'] = camp['_ind'].map({True: 'Sole trader', False: 'Organisation'})
order = {'A': 0, 'B': 1, 'C': 2}
camp['_p'] = camp.Priority.map(order)
camp = camp.sort_values(['_p', '_tier', 'State', 'Company'])
camp = camp.drop_duplicates('Email', keep='first')  # Mission Control also dedupes on the lowercased address
out_cols = ['Email', 'Greeting Name', 'Company', 'Trading / Business Names', 'State', 'Org Size Category', 'Org Type', 'Development Type', 'Priority',
            'Lead Type', 'Website', 'Phone', 'ABN', 'Licence State(s)', 'Master Builders Member', 'Email Status', 'Email Type', 'Email Source', 'Email Shared By (leads)']
valid = camp[camp['Email Status'] == 'Valid'][out_cols]
risky = camp[camp['Email Status'] == 'Risky'][out_cols + ['Email Check']]
valid.to_pickle(f'{W}/v2/campaign_valid.pkl'); risky.to_pickle(f'{W}/v2/campaign_risky.pkl')
valid.to_csv(f'{OUTD}/Mission_Control_Campaign_List_VALID_2026-10-31.csv', index=False, encoding='utf-8')
import xlsxwriter
wb = xlsxwriter.Workbook(f'{OUTD}/Mission_Control_Campaign_List_2026-10-31.xlsx', {'strings_to_urls': False})
H = wb.add_format({'bold': True, 'bg_color': '#1F3864', 'font_color': 'white', 'border': 1})
for name, d in (('Campaign List - Valid', valid), ('Risky - review first', risky)):
    ws = wb.add_worksheet(name)
    for j, c in enumerate(d.columns): ws.write(0, j, c, H); ws.set_column(j, j, 34 if c in ('Email', 'Company', 'Greeting Name', 'Website', 'Email Source', 'Email Check') else 16)
    for i, row in enumerate(d.fillna('').astype(str).values.tolist(), 1):
        for j, v in enumerate(row):
            if v: ws.write_string(i, j, v)
    ws.freeze_panes(1, 1); ws.autofilter(0, 0, len(d), len(d.columns) - 1)
ws = wb.add_worksheet('Read me')
ws.set_column(0, 0, 140)
notes = [
    'HOW TO USE IN MISSION CONTROL',
    'Upload this workbook (or the _VALID.csv) as the campaign list. The importer reads the FIRST sheet ("Campaign List - Valid") and detects the "Email" column by its heading.',
    'Every other column becomes a merge field, e.g. {{greeting_name}}, {{company}}, {{state}}, {{org_size_category}}, {{development_type}}, {{website}}.',
    'Segment or quota by the State column (e.g. one message per state per send time), or by Org Size Category / Priority.',
    '',
    'WHAT "VALID" MEANS HERE',
    '1. The address passes Mission Control\'s own address grammar (src/lib/email/emailAddress.pure.ts, ported).',
    '2. The address was published by the business itself, on its own website (verified by the ABN shown on the site, or by company name plus state), in its association directory listing, or in your own list.',
    '3. Its domain has a live mail server (MX record) right now, checked by DNS on 31 Oct 2026. Domains that do not exist or have no mail server were dropped.',
    '4. Placeholder, system, no-reply, careers/HR, accounts-payable and media addresses were removed. Third-party addresses (e.g. the web designer\'s) were removed.',
    '5. One row per address. The same address listed for several related companies appears once; "Email Shared By (leads)" says how many leads it covers.',
    'NOT checked: whether the individual mailbox exists (SMTP mailbox probing is not possible from this environment). Expect a small bounce rate. Mission Control\'s bounce scan suppresses hard bounces automatically.',
    'Recommended: for a new sending mailbox, follow EMAIL_WARMUP.md and start with the Priority A rows at a low daily cap. Optionally run the list through a mailbox-verification service first.',
    '',
    'RISKY SHEET',
    'Addresses whose domain has no MX record (A-record fallback only), whose DNS check was inconclusive, or whose domain differs from the company website. Review these before adding them to a campaign.',
    '',
    'COMPLIANCE – READ BEFORE SENDING',
    'These addresses were compiled by automated collection from public websites and directories, so under the Spam Act 2003 (Part 3) this is a "harvested-address list".',
    'Using it is lawful only for messages that do not breach section 16, i.e. each recipient has consented. For this list that means INFERRED consent from conspicuous publication (Spam Act Schedule 2, cl. 4).',
    'That requires: (a) the address is conspicuously published (it is: on the business\'s own site or directory listing); (b) the publication has no statement that unsolicited messages are unwanted; (c) your message relates directly to the recipient\'s business or role (e.g. a builder-partnership offer to a builder).',
    'Condition (b) was checked automatically: every page an address came from was re-read, and addresses published next to opt-out wording were removed (see "Do not email" in the master workbook).',
    'Keep messages relevant to building/development businesses, identify the sender accurately, include a working unsubscribe ({{unsubscribe_url}}), and honour it within 5 business days. If in doubt, get legal advice before the first send.',
    'These are business addresses published by the businesses. Under the Spam Act 2003, consent can be inferred from conspicuous publication only where the message relates to the recipient\'s business and the publication does not say unsolicited messages are unwanted.',
    'Every message must identify the sender and carry a working unsubscribe ({{unsubscribe_url}}). Honour unsubscribes within 5 business days.',
]
for i, l in enumerate(notes): ws.write_string(i, 0, l, wb.add_format({'bold': l.isupper() and bool(l)}))
wb.close()
print('campaign valid', len(valid), 'risky', len(risky))
print(valid['Org Size Category'].value_counts().to_dict()); print(valid.State.value_counts().to_dict()); print(valid.Priority.value_counts().to_dict())
# ---- association members whose ABN could not be confirmed: separate, opt-in sheet ----
import importlib.util, sys, requests
sys.argv = ['x']
spec = importlib.util.spec_from_file_location('e2', '/tmp/leads/scripts/enrich_v2.py'); e2 = importlib.util.module_from_spec(spec); spec.loader.exec_module(e2)
ex = pd.read_pickle(f'{W}/excluded_ext.pkl')
mem = ex[(ex.source == 'Master Builders') & (ex.email.fillna('') != '')].copy()
s_ = requests.Session(); s_.headers['accept'] = 'application/dns-json'
def mxs(d):
    try:
        j = s_.get('https://dns.google/resolve', params={'name': d, 'type': 'MX'}, timeout=8).json()
        return 'MX' if j.get('Status') == 0 and any(a.get('type') == 15 for a in j.get('Answer', [])) else 'NXDOMAIN' if j.get('Status') == 3 else 'NO-MX'
    except Exception: return 'DNS-ERROR'
mem['Email'] = mem.email.str.strip().str.lower()
mem = mem[mem.Email.apply(e2.mc_valid) & ~mem.Email.str.contains(e2.JUNK) & ~mem.Email.apply(lambda x: bool(e2.NOTSALES.match(x)))]
mem['Domain mail server'] = mem.Email.str.split('@').str[1].apply(mxs)
mem = mem[mem['Domain mail server'] == 'MX']
mem = mem[~mem.Email.isin(set(valid.Email) | set(risky.Email))].drop_duplicates('Email')
mem = mem.rename(columns={'name': 'Company', 'state': 'State', 'phone': 'Phone', 'website': 'Website', 'reason': 'Why not in main list', 'source_url': 'Directory URL'})
mem['Greeting Name'] = mem.Company
mem = mem[['Email', 'Greeting Name', 'Company', 'State', 'Phone', 'Website', 'Domain mail server', 'Why not in main list', 'Directory URL']]
mem.to_pickle(f'{W}/v2/campaign_members_unconfirmed.pkl')
import openpyxl
wbo = openpyxl.load_workbook(f'{OUTD}/Mission_Control_Campaign_List_2026-10-31.xlsx')
wsn = wbo.create_sheet('Members - ABN unconfirmed', 2)
wsn.append(list(mem.columns))
for row in mem.fillna('').astype(str).values.tolist(): wsn.append(row)
wbo.save(f'{OUTD}/Mission_Control_Campaign_List_2026-10-31.xlsx')
print('association members with ABN unconfirmed (separate sheet)', len(mem))
