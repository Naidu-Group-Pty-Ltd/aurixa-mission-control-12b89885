import duckdb, csv, json, re, os, collections, datetime
import pandas as pd
W = '/tmp/leads/work'; A = '/tmp/leads/agents'
con = duckdb.connect(f'{W}/leads.duckdb', read_only=True)
TODAY = '2026-10-31'
TIERS = ['Enterprise (>$240M)', 'Large ($30M–$240M)', 'Medium ($3M–$30M)', 'Small ($0.8M–$3M)', 'Micro (<$0.8M)']
TIER_IDX = {t: i for i, t in enumerate(TIERS)}
FIN = {'SC1': (4, 'up to $200k'), 'SC2': (4, 'up to $800k'), 'EMRSC1': (4, 'est. up to $200k'), 'EMRSC2': (4, 'est. up to $800k'),
       '1': (3, '$0.8M–$3M'), 'EMR1-2': (3, 'est. $0.8M–$12M'), '2': (2, '$3M–$12M'), '3': (2, '$12M–$30M'), 'EMR3-7': (2, 'est. >$12M'),
       '4': (1, '$30M–$60M'), '5': (1, '$60M–$120M'), '6': (1, '$120M–$240M'), '7': (0, '>$240M')}
FIN_ORDER = ['SC1', 'EMRSC1', 'SC2', 'EMRSC2', '1', 'EMR1-2', '2', '3', 'EMR3-7', '4', '5', '6', '7']
def read(p):
    return list(csv.DictReader(open(p, encoding='utf-8'))) if os.path.exists(p) else []
# ---------------- core register leads (QLD + VIC) ----------------
core = con.execute("select * from lead_core2").df()
core = core.astype(object).where(core.notna(), None)
L = {}
for r in core.to_dict('records'):
    L[r['abn']] = dict(abn=r['abn'], legal_name=r['asic_name'] or r['abr_name'],
        abr_name=r['abr_name'], reg_names={r['reg_name']}, entity_kind=r['entity_kind'], licence_states=set((r['licence_states'] or '').split(', ')) - {''},
        licence_numbers=[r['licence_numbers']], licence_classes=set((r['licence_classes'] or '').split('; ')) - {''},
        qbcc_cats=set((r['qbcc_fin_cats'] or '').replace(' ', '').split(',')) - {''}, address=r['address'] if isinstance(r['address'], str) else '',
        abn_source={r['abn_source']}, sources={('QBCC Licensed Contractors Register' if s == 'QLD' else 'BPC (VBA) Building Practitioner Register') for s in r['licence_states'].split(', ')},
        phone='', email='', website='', web_conf='', contact_source='', mba='', mba_src='', mba_cats='', staff_band='', dev_flags=set(), rank_info='', rank_tier=None,
        orig=False, notes=set(), suburb='', postcode_src='', core_acn=r['acn'])
print('core', len(L))
# ---------------- helper: ABR/ASIC facts for any ABN ----------------
def facts(abns):
    abns = list(abns)
    if not abns: return {}
    con.execute("CREATE OR REPLACE TEMP TABLE q AS SELECT unnest(?) abn", [abns])
    df = con.execute("""select b.abn, b.abn_status, b.abn_status_from, b.entity_type_code, b.entity_type, b.entity_name, b.state, b.postcode, b.gst_status, b.gst_from,
          lpad(nullif(b.acn,''),9,'0') acn, a.asic_name, a.asic_status, a.asic_type, a.asic_class, a.asic_subclass, a.asic_reg_date, x.bns, coalesce(x.n_bn,0) n_bn, b.other_names
        from q join abr b using (abn) left join asic_cur a on a.acn = lpad(nullif(b.acn,''),9,'0') left join bn_abn x on x.abn=b.abn""").df()
    df = df.astype(object).where(df.notna(), None)
    return {r['abn']: r for r in df.to_dict('records')}
# ---------------- external sources ----------------
excluded_ext = []
skipped_ind = {}
def add_external(kind, rows, res, mk, record_excluded=True, require_state=False):
    res = {r['rid']: r for r in res}
    F = facts({r['abn'] for r in res.values() if r['abn']})
    n_new = n_merge = 0
    for i, row in enumerate(rows):
        rr = res.get(str(i), {}); abn = rr.get('abn', '')
        f = F.get(abn) if abn else None
        why = None
        if not abn: why = 'ABN not confirmed: ' + (rr.get('abn_method') or 'no match')
        elif f is None or f['abn_status'] != 'ACT': why = 'ABN not active in ABR'
        elif f['entity_type_code'] in ('PRV', 'PUB') and f['asic_status'] != 'REGD': why = 'ASIC status ' + str(f['asic_status'])
        if why and (rr.get('abn_method') or '').startswith('no unique ABR individual'):
            skipped_ind[kind] = skipped_ind.get(kind, 0) + 1; continue
        if not why and require_state and row.get('state'):
            lic = L[abn]['licence_states'] if abn in L else set()
            if f['state'] != row['state'] and row['state'] not in lic: why = 'matched entity registered in another state'
        if why and not record_excluded: continue
        if why:
            excluded_ext.append(dict(source=kind, name=row.get('business_name') or row.get('company_name') or row.get('entity_name'), state=row.get('state', ''),
                                     reason=why, phone=row.get('phone', ''), email=row.get('email', ''), website=row.get('website', ''), source_url=row.get('source_url', ''))); continue
        if abn in L: n_merge += 1
        else:
            n_new += 1
            L[abn] = dict(abn=abn, legal_name=f['asic_name'] if isinstance(f['asic_name'], str) else f['entity_name'], abr_name=f['entity_name'], reg_names=set(),
                entity_kind='Individual' if f['entity_type_code'] == 'IND' else 'Company', licence_states=set(), licence_numbers=[], licence_classes=set(), qbcc_cats=set(),
                address='', abn_source=set(), sources=set(), phone='', email='', website='', web_conf='', contact_source='', mba='', mba_src='', mba_cats='', staff_band='',
                dev_flags=set(), rank_info='', rank_tier=None, orig=False, notes=set(), suburb='', postcode_src='')
        L[abn]['abn_source'].add(rr['abn_method'])
        mk(L[abn], row)
    print(kind, 'merged', n_merge, 'new', n_new, 'excluded', sum(1 for e in excluded_ext if e['source'] == kind))
def set_contact(d, row, label):
    for k in ('phone', 'email', 'website'):
        v = (row.get(k) or '').strip()
        if v and not d[k]: d[k] = v; d['contact_source'] = label
    if not d['address'] and row.get('address'):
        d['address'] = ', '.join(x for x in [row.get('address', ''), row.get('suburb', ''), row.get('state', ''), row.get('postcode', '')] if x)
    if not d['suburb'] and row.get('suburb'): d['suburb'] = row['suburb']
# other-state registers
regs = read(f'{W}/regs_filtered.csv')
if regs and os.path.exists(f'{W}/regs_resolved.csv'):
    def mk_reg(d, row):
        d['licence_states'].add(row['state'].strip()); d['licence_numbers'].append(f"{row['state']}: {row['licence_number']}")
        if row.get('licence_class'): d['licence_classes'].add(row['licence_class'])
        d['sources'].add(row['register']); d['reg_names'].add(row['entity_name']); set_contact(d, row, row['register'])
        if 'not re-verified' in row.get('licence_status', ''): d['notes'].add('NSW licence taken from the Fair Trading contractor-licence file (Dec 2023), expiry after 31 Oct 2026, not live-checked')
    add_external('State registers (NSW/WA/SA/TAS/ACT/NT)', regs, read(f'{W}/regs_resolved.csv'), mk_reg)
# Master Builders
mba = read(f'{W}/mba_filtered.csv')
if mba and os.path.exists(f'{W}/mba_resolved.csv'):
    def mk_mba(d, row):
        d['mba'] = 'Yes'; d['mba_src'] = '; '.join(sorted(set(filter(None, d['mba_src'].split('; ') + [row['source']]))))
        d['sources'].add(row['source'] + ' member directory')
        if row.get('member_categories') and not d['mba_cats']: d['mba_cats'] = row['member_categories'][:300]
        m = re.search(r'Size:\s*([^;]+?Staff)', row.get('notes', '') or '')
        if m and not d['staff_band']: d['staff_band'] = m.group(1).strip()
        set_contact(d, row, row['source'] + ' directory')
        if row.get('trading_name'): d['reg_names'].add(row['trading_name'])
        if row.get('business_name'): d['reg_names'].add(row['business_name'])
    add_external('Master Builders', mba, read(f'{W}/mba_resolved.csv'), mk_mba)
# Developers / rankings / panels
dev = read(f'{A}/dev/developers_and_rankings.csv')
def rank_tier(row):
    v = (row.get('size_indicator_value') or '')
    m = re.search(r'national rank (?:equal )?(\d+)', v); starts = re.search(r'([\d,]+) starts', v)
    n = int(m.group(1)) if m else None; s = int(starts.group(1).replace(',', '')) if starts else None
    if (n and n <= 20) or (s and s >= 1000) or 'group 8,038' in v: return 0
    if 'Housing 100' in row['source'] or 'Housing 100' in (row.get('notes') or ''): return 1
    return None
if dev:
    def mk_dev(d, row):
        t = (row.get('org_type') or '').lower()
        if 'develop' in t: d['dev_flags'].add('Developer')
        src = row['source']; d['sources'].add(src)
        rt = rank_tier(row)
        if 'Housing 100' in src or 'Housing 100' in (row.get('notes') or ''):
            d['rank_info'] = '; '.join(filter(None, [d['rank_info'], 'HIA Housing 100: ' + row['size_indicator_value'][:120]]))
        elif 'Panel' in src or 'Prequalification' in src:
            d['notes'].add(src)
        if rt is not None and (d['rank_tier'] is None or rt < d['rank_tier']): d['rank_tier'] = rt
        set_contact(d, row, src)
    add_external('Developers & rankings', dev, read(f'{W}/dev_resolved.csv'), mk_dev)
# the client's own Top-10 list: validated entries join the main list (contact details and call notes carried over)
import openpyxl as _ox
_all = list(_ox.load_workbook('/tmp/input_xlsx/input.xlsx').worksheets[0].iter_rows(min_row=2, values_only=True))
_mm = {r['rid']: r for r in read(f'{W}/orig_match.csv')}
_compact = [(i, r) for i, r in enumerate(_all) if r[1]]
orig_rows = [dict(business_name=r[1], state=r[0] or '', phone=str(r[4] or ''), email=r[5] or '', website=r[6] or '', notes=(r[7] or '').strip()) for i, r in _compact]
orig_res = [dict(rid=str(j), abn=_mm.get(str(i), {}).get('abn', ''), abn_method=_mm.get(str(i), {}).get('abn_match_method', '')) for j, (i, r) in enumerate(_compact)]
def mk_orig(d, row):
    d['orig'] = True; d['sources'].add('Your Top-10 builders/developers list (contact details as supplied)')
    set_contact(d, row, 'Your Top-10 list')
    if row['notes']: d['call_notes'] = (d.get('call_notes', '') + ' ' + row['notes']).strip()
    d.setdefault('grp_ev', []).append((1, 'Listed in your Top-10 builders/developers list (%s)' % row['business_name'], 'Medium', ''))
add_external('Your Top-10 list', orig_rows, orig_res, mk_orig, record_excluded=False, require_state=True)
# original list flag
for m in read(f'{W}/orig_match.csv'):
    if m['abn'] and m['abn'] in L: L[m['abn']]['orig'] = True
# web enrichment
import itertools
for l in itertools.chain(open(f'{W}/web/web_out.jsonl'), open(f'{W}/web/web_out2.jsonl') if os.path.exists(f'{W}/web/web_out2.jsonl') else []):
    j = json.loads(l)
    d = L.get(j['abn'])
    if not d or not j.get('website'): continue
    if not d['website']: d['website'] = j['website']; d['web_conf'] = j['web_conf']
    if not d['email'] and j.get('emails'): d['email'] = j['emails'][0]
    if not d['phone'] and j.get('phones'): d['phone'] = j['phones'][0]
    if not d['contact_source']: d['contact_source'] = 'Company website (' + j['web_conf'] + ')'
    elif 'website' not in d['contact_source']: d['contact_source'] += '; company website'
# ---------------- group-level evidence by distinctive name ----------------
STOP = {'THE', 'TRUSTEE', 'FOR', 'ATF', 'PTY', 'PROPRIETARY', 'LTD', 'LIMITED', 'CO', 'AND', 'TRUST', 'UNIT', 'FAMILY', 'DISCRETIONARY', 'P', 'L', 'AS'}
STATES = {'QLD', 'VIC', 'NSW', 'WA', 'SA', 'TAS', 'ACT', 'NT', 'AUST', 'AUSTRALIA', 'HOLDINGS'}
def pkey(x):
    t = [w for w in re.sub(r'[^A-Z0-9 ]', ' ', (x or '').upper()).split() if w not in STOP]
    while t and t[-1] in STATES: t.pop()
    return ' '.join(t)
grp = {}
for row in dev:
    rt = rank_tier(row)
    if rt is None: continue
    k = pkey(re.sub(r'\(.*?\)', ' ', row['company_name']))
    if len(k.split()) >= 2: grp.setdefault(k, []).append((rt, 'HIA Housing 100 ranked builder (group/brand: %s)' % row['company_name'], 'High', row.get('state', '')))
import openpyxl
_ws = openpyxl.load_workbook('/tmp/input_xlsx/input.xlsx').worksheets[0]
orig_keys = set()
for r_ in _ws.iter_rows(min_row=2, values_only=True):
    if r_[1]:
        k = pkey(re.sub(r'\(.*?\)', ' ', r_[1]))
        if len(k.split()) >= 2:
            orig_keys.add((k, r_[0])); grp.setdefault(k, []).append((1, "Listed in your Top-10 builders/developers list (%s)" % r_[1], 'Medium', r_[0] or ''))
n_grp = 0
for abn, d in L.items():
    keys = {pkey(d['legal_name']), pkey(d['abr_name'])} | {pkey(n) for n in d['reg_names'] if n}
    sts = set(d['licence_states'])
    for k in keys:
        for ev in grp.get(k, []):
            if ev[3] and sts and ev[3] not in sts: continue
            if ev[3] and not sts: continue
            d.setdefault('grp_ev', []).append(ev); n_grp += 1
            if ev[2] == 'Medium': d['orig'] = True
print('group-level size evidence applied', n_grp)
# ---------------- facts for everything ----------------
F = facts(L.keys())
_need = {d['core_acn']: abn for abn, d in L.items() if d.get('core_acn') and abn in F and not F[abn]['asic_status']}
if _need:
    con.execute("CREATE OR REPLACE TEMP TABLE qa AS SELECT unnest(?) acn", [list(_need)])
    for acn_, nm_, st_, ty_, rd_ in con.execute("select a.acn, asic_name, asic_status, asic_type, asic_reg_date from qa join asic_cur a using (acn)").fetchall():
        f_ = F[_need[acn_]]; f_['asic_status'] = st_; f_['asic_type'] = ty_; f_['asic_reg_date'] = rd_; f_['trustee_acn'] = acn_; f_['trustee_name'] = nm_
print('trustee ACN facts added', len(_need))
DEVRE = re.compile(r'\b(DEVELOPMENTS?|DEVELOPERS?|PROPERTY|PROPERTIES|LAND|ESTATES?|COMMUNITIES|URBAN)\b', re.I)
HOMERE = re.compile(r'\bHOMES?\b', re.I)
TRADE = re.compile(r'(PLUMB|ELECTRIC|ROOF|SOLAR|ENERGY|ENGINEER|LANDSCAP|POOL|PAINT|CONCRET|STEEL|SHED|MEZZANINE|FLOOR|GLASS|GLAZ|KITCHEN|CABINET|JOINERY|BITUMEN|CIVIL|FENC|SCAFFOLD|DEMOLI|EXCAVAT|FIRE|AIR ?CON|REFRIG|MECHANICAL|SHOPFIT|INTERIOR|TILING|WATERPROOF|CARPENTRY|FRAMING|BRICK|PLASTER|CEILING|INSULAT|SIGN|MARINE|PEST|CLEAN|PATIO|AWNING|BLIND|DOOR|WINDOW|CRANE|RIGGING|DRILL|PILING|DRAIN)', re.I)
CORE = re.compile(r'(HOME|BUILD|CONSTRUCT|DEVELOP|PROPERT|PROJECT|HOUSING|RESIDENTIAL|DWELLING|ESTATE|COMMUNIT)', re.I)
def staff_tier(b):
    nums = [int(x) for x in re.findall(r'\d+', b or '')]
    if not nums: return None
    top = max(nums)
    return 4 if top <= 5 else 3 if top <= 20 else 2 if top <= 100 else 1 if top <= 500 else 0
def fmt_phone(v):
    for part in re.split(r'[;,/]| or ', v or ''):
        p = re.sub(r'\D', '', part)
        if p.startswith('61') and len(p) == 11: p = '0' + p[2:]
        if len(p) == 9 and p[0] in '23478': p = '0' + p
        if len(p) == 10 and p[:2] == '04': return f'{p[:4]} {p[4:7]} {p[7:]}'
        if len(p) == 10 and p[:4] in ('1300', '1800'): return f'{p[:4]} {p[4:7]} {p[7:]}'
        if len(p) == 10 and p[0] == '0': return f'{p[:2]} {p[2:6]} {p[6:]}'
        if len(p) == 6 and p[:2] == '13': return f'{p[:2]} {p[2:4]} {p[4:]}'
    return (v or '').strip()
out = []
for abn, d in L.items():
    f = F.get(abn)
    if f is None: continue
    age = None
    try: age = 2026 - int(str(f['asic_reg_date'])[-4:])
    except Exception:
        try: age = 2026 - int(str(f['abn_status_from'])[:4])
        except Exception: pass
    sigs = []; tier = None; basis = ''; conf = ''
    cats = sorted(d['qbcc_cats'], key=lambda c: FIN_ORDER.index(c) if c in FIN_ORDER else -1)
    high = []
    if cats and cats[-1] in FIN:
        t, band = FIN[cats[-1]]; high.append((t, f'QBCC financial category {cats[-1]} (allowable annual revenue {band})'))
    if d['rank_tier'] is not None: high.append((d['rank_tier'], 'HIA Housing 100 ranked builder'))
    for t_, b_, c_, _s in d.get('grp_ev', []):
        if c_ == 'High': high.append((t_, b_))
    if f['asic_subclass'] in ('LIST', 'LISN'): high.append((0, 'ASIC listed public company'))
    if high:
        tier, basis = min(high); conf = 'High' if not cats or not cats[-1].startswith('EMR') or len(high) > 1 else 'Medium'
        if len(high) > 1: basis = '; '.join(b for _, b in sorted(high))
    elif any(e[2] == 'Medium' for e in d.get('grp_ev', [])):
        tier, basis = min((e[0], e[1]) for e in d['grp_ev'] if e[2] == 'Medium'); conf = 'Medium'
    elif staff_tier(d['staff_band']) is not None:
        tier = staff_tier(d['staff_band']); basis = f"Master Builders directory staff band: {d['staff_band']}"; conf = 'Medium'
    else:
        sc = 0
        if f['asic_type'] == 'APUB': sc += 4; sigs.append('public company')
        if len(d['licence_states']) >= 2: sc += 2; sigs.append('licensed in %d states' % len(d['licence_states']))
        cl = ' '.join(d['licence_classes'])
        if any(k in cl for k in ('Commercial Builder - Unlimited', 'Builder - Open', 'Builder - Class A', 'Commercial (Unrestricted)')): sc += 2; sigs.append('unrestricted commercial class')
        elif any(k in cl for k in ('Commercial Builder - Limited', 'Medium Rise', 'Builder - Class B', 'Commercial (Restricted)')): sc += 1; sigs.append('commercial/medium-rise class')
        if age is not None and age >= 25: sc += 2; sigs.append(f'{age} yrs registered')
        elif age is not None and age >= 10: sc += 1; sigs.append(f'{age} yrs registered')
        if (f['n_bn'] or 0) >= 4: sc += 2; sigs.append(f"{f['n_bn']} business names")
        elif (f['n_bn'] or 0) >= 2: sc += 1; sigs.append(f"{f['n_bn']} business names")
        if d['mba']: sigs.append('Master Builders member')
        if f['gst_status'] != 'ACT': tier = 4; sigs = ['not GST-registered (turnover likely <$75k)']
        elif f['entity_type_code'] == 'IND': tier = 4; sigs.append('sole trader')
        else: tier = 4 if sc == 0 else 3 if sc <= 2 else 2 if sc <= 5 else 1 if sc <= 7 else 0
        basis = 'Estimated from public signals (score %d): %s' % (sc, ', '.join(sigs) or 'no size signals; young single-class company'); conf = 'Low'
    # org type / development type
    names = ' '.join([d['legal_name'] or ''] + list(d['reg_names']) + [f['bns'] or ''])
    is_dev = 'Developer' in d['dev_flags'] or bool(DEVRE.search(names))
    licensed = bool(d['licence_states'])
    org_type = ('Builder & Developer' if licensed else 'Developer') if is_dev else 'Builder'
    lname = d['legal_name'] or ''
    specialist = (not is_dev) and bool(TRADE.search(lname)) and not CORE.search(lname) and not d['rank_info']
    if specialist: org_type = 'Specialist contractor (holds builder licence)'
    cl = d['licence_classes']; dt = []
    if any('Domestic Builder' in c for c in cl) or 'Builder - Low Rise' in cl or HOMERE.search(names): dt.append('Residential')
    if any('Commercial' in c for c in cl) or 'Builder - Open' in cl or any('Class A' in c for c in cl): dt.append('Commercial')
    if any('Residential' in c for c in cl) or any('Class C' in c for c in cl): dt.append('Residential')
    if 'Builder - Medium Rise' in cl or 'Builder - Open' in cl: dt.append('Multi-unit / Medium-High Rise')
    if any('Project Manage' in c for c in cl): dt.append('Project Management')
    if d['rank_info']: dt.insert(0, 'Volume Housing')
    if is_dev: dt.append('Property Development')
    for c in cl:
        if c and not any(k in c.lower() for k in ('builder', 'project', 'building contractor')): dt.append(c[:60])
    dev_type = ' / '.join(dict.fromkeys(dt)) or 'Building'
    lic_txt = '; '.join(sorted(cl))
    desc = []
    if d['licence_states']: desc.append('Licensed builder in ' + ', '.join(sorted(d['licence_states'])))
    if cats: desc.append('QBCC financial category ' + cats[-1])
    if d['rank_info']: desc.append(d['rank_info'])
    if d['mba']: desc.append('Master Builders member (' + d['mba_src'].replace('Master Builders ', '') + ')')
    if d['notes']: desc.append('; '.join(sorted(d['notes'])))
    if 'Developer' in d['dev_flags']: desc.append('Listed as developer by industry body')
    tn = [x for x in (f['bns'] or '').split(' | ') if x][:4]
    has_contact = bool(d['phone'] or d['email'])
    ent = f['entity_type'] or ''
    is_ind = f['entity_type_code'] == 'IND'
    tier_name = TIERS[tier]
    prio = 'A' if tier <= 2 and has_contact else 'B' if (tier <= 2 or (tier == 3 and has_contact) or d['mba']) else 'C'
    if specialist and tier > 1: prio = 'C'
    pc = f['postcode'] or ''
    out.append({
        'Priority': prio,
        'State': f['state'] or (sorted(d['licence_states'])[0] if d['licence_states'] else ''),
        'Company': (d['legal_name'] or '').strip(),
        'Trading / Business Names': ' | '.join(tn),
        'Org Type': org_type,
        'Development Type': dev_type,
        'Org Size Category': tier_name,
        'Size Basis': basis,
        'Size Confidence': conf,
        'Description': '. '.join(desc),
        'Phone': fmt_phone(d['phone']), 'Email': (d['email'] or '').strip().lower(), 'Website': (d['website'] or '').strip(),
        'Contact Source': d['contact_source'],
        'Website Match': d['web_conf'],
        'Business Address': '' if is_ind else (d['address'] or ''),
        'Postcode (ABR)': pc,
        'ABN': f"{abn[:2]} {abn[2:5]} {abn[5:8]} {abn[8:]}",
        'ABN Status (ABR)': 'Active' if f['abn_status'] == 'ACT' else f['abn_status'],
        'ABN Active Since': f"{str(f['abn_status_from'])[6:8]}/{str(f['abn_status_from'])[4:6]}/{str(f['abn_status_from'])[:4]}" if f['abn_status_from'] else '',
        'Entity Type (ABR)': ent,
        'GST Registered': 'Yes' if f['gst_status'] == 'ACT' else 'No',
        'ACN': (f"{f['acn'][:3]} {f['acn'][3:6]} {f['acn'][6:]}" if isinstance(f['acn'], str) and f['acn'].strip('0') else '') or (f"{f['trustee_acn'][:3]} {f['trustee_acn'][3:6]} {f['trustee_acn'][6:]} (trustee company)" if f.get('trustee_acn') else ''),
        'ASIC Status': {'REGD': 'Registered'}.get(f['asic_status'], f['asic_status'] or ('n/a (not a company)' if f['entity_type_code'] not in ('PRV', 'PUB') else '')),
        'ASIC Company Type': {'APTY': 'Australian Proprietary', 'APUB': 'Australian Public'}.get(f['asic_type'], f['asic_type'] or ''),
        'ASIC Registered': f['asic_reg_date'] or '',
        'ABN Resolved Via': '; '.join(sorted(d['abn_source'])),
        'Validation': 'PASS: ABN checksum valid; ABN active in ABR (extract 28 Oct 2026)' + ('; ASIC registered (dataset Oct 2026)' if f['asic_status'] == 'REGD' else '') + ('; licence current on state register' if d['licence_states'] else '') + ('; NOTE: NSW licence from Dec-2023 file' if any('Dec 2023' in n for n in d['notes']) else ''),
        'Licence State(s)': ', '.join(sorted(d['licence_states'])),
        'Licence Number(s)': '; '.join(d['licence_numbers'])[:250],
        'Licence Class(es)': lic_txt,
        'QBCC Financial Category': ', '.join(cats),
        'Master Builders Member': d['mba'] or '',
        'MBA Directory Categories': d['mba_cats'],
        'Staff Band (MBA)': d['staff_band'],
        'In Original List': 'Yes' if d['orig'] else '',
        'Sources': '; '.join(sorted(d['sources'])),
        'ABN Lookup': f'https://abr.business.gov.au/ABN/View?abn={abn}',
        'Notes': d.get('call_notes', ''),
        '_tier': tier, '_ind': is_ind, '_contact': has_contact,
    })
df = pd.DataFrame(out)
df = df.sort_values(['_tier', '_contact', 'Master Builders Member', 'State', 'Company'], ascending=[True, False, False, True, True])
df.to_pickle(f'{W}/final_df.pkl')
pd.DataFrame(excluded_ext).to_pickle(f'{W}/excluded_ext.pkl')
print(len(df), df['_ind'].sum(), 'unmatched individuals not listed:', skipped_ind)
json.dump(skipped_ind, open(f'{W}/skipped_ind.json', 'w'))
print(df.groupby(['Org Size Category', 'Size Confidence']).size())
print(df.groupby('State').size())
print('with phone', (df.Phone != '').sum(), 'email', (df.Email != '').sum(), 'web', (df.Website != '').sum())
